"""
Build the in-game test workbook for the v0.136 tester baseline (docs/V0136_ABILITY_TEST_PLAN.xlsx).

One row per ability the baseline touched (every AbilityMap entry carrying the `[v0.136 baseline]` note):
  • "Presets"       — enabled abilities with a pre-set fix: what changed, the value, native values, tester
                      feedback, the expected result, the exact commands to test it, + Actual/Worked? to fill in.
  • "Hard-blocked"  — code-locked crash/character-break abilities: confirm they can't be captured.
  • "Soft-disabled" — switched off by default: spot-check they don't drop, and pick Keep off / Re-enable.
Reads (read-only): Resources/ability_rules.default.json, ability_metadata.json, prefab_names.tsv, the tester
matrix xlsx (for the full tester finding), tools/apply_critical_blocks.py (HARDBLOCK set) and
docs/INGAME_TEST_CHECKLIST.md §7 (rows there are marked High priority).
Refuses to overwrite an existing workbook (it holds your results) unless --force.
"""
import os, re, sys, json, math, argparse, datetime
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.join(HERE, '..', 'Beelzebub')
RES = os.path.join(PROJ, 'Resources')
OUT = os.path.join(PROJ, 'docs', 'V0136_ABILITY_TEST_PLAN.xlsx')
MATRIX = os.path.join(HERE, '..', '..', 'Reference Data', 'TESTING DATA', 'TESTER_ABILITY_MATRIX.xlsx')
CHECKLIST = os.path.join(PROJ, 'docs', 'INGAME_TEST_CHECKLIST.md')
TAG = '[v0.136 baseline]'

KNOBS = [  # (json key, admin knob name, label)
    ('FreeMoveAfterSeconds', 'freelymove', 'Unrooted'),
    ('CooldownSeconds', 'cooldown', 'Cooldown'),
    ('CooldownScale', 'cooldownscale', 'Cooldown'),
    ('DamageScale', 'damagescale', 'Damage'),
    ('LeapHeight', 'leapheight', 'Leap clamp'),
    ('SummonCap', 'summoncap', 'Summon limit'),
    ('SummonTimeoutSeconds', 'summontimeout', 'Summon limit'),
    ('MaxRangeOverride', 'range', 'Range'),
    ('ForceTimeoutSeconds', 'forcetimeout', 'Effect timeout'),
    ('Weapons', 'weapons', 'Weapon limit'),
]
RESULTS = ['Pass', 'Partial', 'Fail', "Can't test"]

# ── palette ──
INK, MUTED = '1F2328', '57606A'
HEAD_FILL = PatternFill('solid', fgColor='3B1F2B')          # dark wine
INPUT_FILL = PatternFill('solid', fgColor='FFF8E1')         # soft yellow = "you fill this in"
KEY_FILL = PatternFill('solid', fgColor='F3EEF1')
PASS_FILL = PatternFill('solid', fgColor='D4EDDA')
PART_FILL = PatternFill('solid', fgColor='FFF3CD')
FAIL_FILL = PatternFill('solid', fgColor='F8D7DA')
SKIP_FILL = PatternFill('solid', fgColor='E2E3E5')
HIGH_FONT = Font(bold=True, color='9A1B3A')
THIN = Side(style='thin', color='D0D7DE')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='top', wrap_text=True)


def fmt(v):
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    if isinstance(v, list): return ', '.join(map(str, v))
    return str(v)


def load():
    rules = json.load(open(os.path.join(RES, 'ability_rules.default.json'), encoding='utf-8'))['AbilityMap']
    meta = json.load(open(os.path.join(RES, 'ability_metadata.json'), encoding='utf-8'))['abilities']
    n2g = {}
    for line in open(os.path.join(RES, 'prefab_names.tsv'), encoding='utf-8'):
        p = line.rstrip('\n').split('\t')
        if len(p) >= 2:
            try: n2g.setdefault(p[1], int(p[0]))
            except ValueError: pass
    findings = {}
    if os.path.exists(MATRIX):
        rows = list(openpyxl.load_workbook(MATRIX, read_only=True).worksheets[0].iter_rows(values_only=True))
        ix = {h: i for i, h in enumerate(rows[0])}
        for r in rows[1:]:
            try: gid = int(r[ix['ID']])
            except (TypeError, ValueError): continue
            f = r[ix['Phase 2: Finding']] or r[ix['Phase 1: TesterFinding']] or r[ix['Phase 1: Notes']]
            if f: findings[gid] = str(f).strip()
    src = open(os.path.join(HERE, 'apply_critical_blocks.py'), encoding='utf-8').read()
    hard = {int(x) for x in re.findall(r'-?\d{5,}', src[src.index('HARDBLOCK = {'):src.index('}', src.index('HARDBLOCK = {'))])}
    ck = open(CHECKLIST, encoding='utf-8').read()
    ck = ck[ck.find('## 7.'):]
    priority = {int(b) for a, b in re.findall(r'`(-?\d+) (-?\d+)`', ck)}
    return rules, meta, n2g, findings, hard, priority


def tester_text(note, gid, findings):
    t = findings.get(gid)
    if not t:
        i = (note or '').find('Tester:')
        t = note[i + 7:].strip() if i >= 0 else ''
    t = re.sub(r'\s+', ' ', t).strip()
    return t[:600] + ('…' if len(t) > 600 else '')


def expected(e, base_cd, base_ct, reenabled):
    out = []
    if reenabled: out.append('Can be captured/granted again (it was switched off in the first pass).')
    for key, knob, _ in KNOBS:
        if key not in e: continue
        v = e[key]
        if key == 'FreeMoveAfterSeconds':
            out.append(f'You can walk about {fmt(v)}s into the cast instead of standing rooted for the whole '
                       f'cast/channel. The ability still fires and hits normally.')
        elif key == 'CooldownSeconds':
            if not base_cd or base_cd <= 0.01:
                out.append(f"Can't recast for {fmt(v)}s. It had no cooldown of its own, so Beelzebub starts one: "
                           f'the client may not show a cooldown swipe, but the recast must be refused.')
            else:
                out.append(f"Can't recast for {fmt(v)}s (native {fmt(base_cd)}s).")
        elif key == 'CooldownScale':
            eff = f'about {fmt(round(base_cd * v, 1))}s' if base_cd else f'{fmt(v)}× the native cooldown'
            out.append(f'Cooldown is {eff} (native {fmt(base_cd or 0)}s × {fmt(v)}).')
        elif key == 'DamageScale':
            word = 'weaker' if v < 1 else 'stronger'
            out.append(f'Noticeably {word}: about {fmt(v)}× the damage it did before (compare on a training dummy).')
        elif key == 'LeapHeight':
            out.append('A short, low hop instead of the big launch into the sky; you land safely and can act.')
        elif key == 'SummonCap':
            out.append(f'At most {fmt(v)} summon alive at once; recasting does not add more.')
        elif key == 'SummonTimeoutSeconds':
            out.append(f'Summons despawn on their own after about {fmt(v)}s.')
        elif key == 'MaxRangeOverride':
            out.append(f"Can't be aimed/thrown farther than about {fmt(v)}m.")
        elif key == 'ForceTimeoutSeconds':
            out.append(f'The effect ends on its own after about {fmt(v)}s — you never get stuck in it.')
        elif key == 'Weapons':
            out.append(f'Only usable with: {fmt(v)}.')
    return '\n'.join('• ' + s for s in out)


def unit_of(meta_e):
    for n in meta_e.get('sourceNpcs') or []:
        if n.get('name') and 'Blood Soul' not in n['name']:
            return n['guid'], n['name']
    s = (meta_e.get('sourceNpcs') or [None])[0]
    return (s['guid'], s.get('name') or '') if s else (0, '')


def build():
    rules, meta, n2g, findings, hard_guids, priority = load()
    presets, hard, soft = [], [], []
    for name, e in rules.items():
        note = e.get('Notes') or ''
        gid = n2g.get(name, 0)
        if TAG not in note and gid not in hard_guids: continue
        m = meta.get(str(gid), {})
        ugid, uname = unit_of(m)
        disp = m.get('name') or name
        knobs = [(knob, e[key]) for key, knob, _ in KNOBS if key in e]
        types = []
        for key, _, label in KNOBS:
            if key in e and label not in types: types.append(label)
        row = dict(name=name, gid=gid, disp=disp, unit=uname or '—', ugid=ugid,
                   knobs='\n'.join(f'{k} = {fmt(v)}' for k, v in knobs),
                   native=f"cooldown {fmt(m.get('baseCooldown', '?'))}s · cast {fmt(m.get('baseCastTime', '?'))}s",
                   tester=tester_text(note, gid, findings), note=note[note.find(TAG) + len(TAG):].strip() if TAG in note else note)
        give = f'.beelz admin give <you> {ugid} {gid}\n.beelz grant <slot> {gid}'
        check = f'.beelz admin ability {name}'
        blocked = e.get('Enabled') is False or e.get('ReviewStatus') in ('Blocked', 'Hidden')
        if gid in hard_guids or 'HARD-BLOCKED' in note:
            row.update(why=row['tester'] or row['note'], cmd=check,
                       exp="Never drops from the unit and can't be granted into use. "
                           f'`{check}` shows enabled=off, reviewstatus=Blocked.')
            hard.append(row)
        elif blocked:
            row.update(cmd=check, ready=row['knobs'] or '—',
                       exp="Doesn't drop from this unit by default. "
                           f'`{check}` shows enabled=off. Re-enable with `{check} enabled true`.')
            soft.append(row)
        elif knobs:
            row.update(types=' + '.join(types), cmd=f'{give}\nCheck values: {check}',
                       exp=expected(e, m.get('baseCooldown'), m.get('baseCastTime'), note.startswith('Re-audit: enabled')),
                       prio='High' if gid in priority or (
                           'CooldownSeconds' in e and (m.get('baseCooldown') or 0) <= 0.01) else 'Normal')
            presets.append(row)
    presets.sort(key=lambda r: (r['prio'] != 'High', r['types'], r['unit'], r['disp']))
    hard.sort(key=lambda r: (r['unit'], r['disp']))
    soft.sort(key=lambda r: (r['unit'], r['disp']))
    return presets, hard, soft


def height(values, widths):
    lines = 1
    for v, w in zip(values, widths):
        s = str(v or '')
        n = sum(max(1, math.ceil(len(part) / max(1, w * 1.15))) for part in s.split('\n'))
        lines = max(lines, n)
    return min(15 * lines + 4, 300)


def sheet(wb, title, cols, rows, tab_color, table_name, result_col=None, extra_dv=None):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab_color
    heads = [c[0] for c in cols]
    widths = [c[1] for c in cols]
    ws.append(heads)
    for r in rows: ws.append([c[2](r) for c in cols])
    last = len(rows) + 1
    for i, (h, w, _, kind) in enumerate(cols, 1):
        L = get_column_letter(i)
        ws.column_dimensions[L].width = w
        hc = ws.cell(1, i)
        hc.font = Font(bold=True, color='FFFFFF'); hc.fill = HEAD_FILL; hc.alignment = CENTER; hc.border = BOX
        for rr in range(2, last + 1):
            c = ws.cell(rr, i)
            c.border = BOX
            c.alignment = CENTER if kind in ('center', 'result', 'choice') else WRAP
            if kind in ('result', 'input', 'choice'): c.fill = INPUT_FILL
            if kind == 'mono': c.font = Font(name='Consolas', size=9, color=INK)
            if kind == 'key': c.fill = KEY_FILL; c.font = Font(bold=True, color=INK)
            if kind == 'muted': c.font = Font(size=9, color=MUTED)
    ws.row_dimensions[1].height = 32
    for rr in range(2, last + 1):
        ws.row_dimensions[rr].height = height([ws.cell(rr, i).value for i in range(1, len(cols) + 1)], widths)
    t = Table(displayName=table_name, ref=f'A1:{get_column_letter(len(cols))}{last}')
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
    ws.add_table(t)
    ws.freeze_panes = ws.cell(2, 5 if len(cols) > 8 else 3)
    for i, (_, _, _, kind) in enumerate(cols, 1):
        L = get_column_letter(i)
        rng = f'{L}2:{L}{last}'
        if kind == 'result':
            dv = DataValidation(type='list', formula1='"' + ','.join(RESULTS) + '"', allow_blank=True)
            dv.error = 'Pick Pass, Partial, Fail or Can\'t test'; dv.prompt = 'Did it behave as expected?'
            ws.add_data_validation(dv); dv.add(rng)
            for val, fill in (('Pass', PASS_FILL), ('Partial', PART_FILL), ('Fail', FAIL_FILL), ("Can't test", SKIP_FILL)):
                ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=[f'"{val}"'], fill=fill))
        if kind == 'choice' and extra_dv:
            dv = DataValidation(type='list', formula1='"' + ','.join(extra_dv) + '"', allow_blank=True)
            ws.add_data_validation(dv); dv.add(rng)
        if kind == 'prio':
            ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['"High"'], font=HIGH_FONT))
    ws.sheet_view.zoomScale = 90
    return ws


def overview(wb, presets, hard, soft):
    ws = wb.active
    ws.title = 'Start here'
    ws.sheet_properties.tabColor = '3B1F2B'
    ws.column_dimensions['A'].width = 30
    for L in 'BCDEFG': ws.column_dimensions[L].width = 13
    ws['A1'] = 'Beelzebub v0.136 — ability test plan'
    ws['A1'].font = Font(size=18, bold=True, color='3B1F2B')
    ws['A2'] = (f'Generated {datetime.date.today():%Y-%m-%d} from the shipped default rules. Every ability the Discord '
                'tester baseline changed is listed. Fill in the yellow columns.')
    ws['A2'].font = Font(italic=True, color=MUTED)
    ws.merge_cells('A2:G2')
    ws['A2'].alignment = WRAP
    ws.row_dimensions[2].height = 32

    r = 4
    ws.cell(r, 1, 'Progress').font = Font(size=13, bold=True, color='3B1F2B')
    r += 1
    heads = ['Sheet', 'Abilities', 'Pass', 'Partial', 'Fail', "Can't test", 'Done %']
    for i, h in enumerate(heads, 1):
        c = ws.cell(r, i, h); c.font = Font(bold=True, color='FFFFFF'); c.fill = HEAD_FILL; c.alignment = CENTER; c.border = BOX
    specs = [('Presets', 'Presets', len(presets), 'L'), ('Hard-blocked', "'Hard-blocked'", len(hard), 'H'),
             ('Soft-disabled', "'Soft-disabled'", len(soft), 'H')]
    first = r + 1
    for title, ref, n, col in specs:
        r += 1
        rng = f'{ref}!{col}2:{col}{n + 1}'
        vals = [title, n] + [f'=COUNTIF({rng},"{v}")' for v in RESULTS] + [f'=IF(B{r}=0,0,SUM(C{r}:F{r})/B{r})']
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v); c.border = BOX; c.alignment = CENTER if i > 1 else WRAP
        ws.cell(r, 7).number_format = '0%'
    r += 1
    ws.cell(r, 1, 'Total').font = Font(bold=True)
    for i in range(2, 7):
        L = get_column_letter(i)
        ws.cell(r, i, f'=SUM({L}{first}:{L}{r - 1})').font = Font(bold=True)
    ws.cell(r, 7, f'=IF(B{r}=0,0,SUM(C{r}:F{r})/B{r})').number_format = '0%'
    for i in range(1, 8):
        ws.cell(r, i).border = BOX; ws.cell(r, i).alignment = CENTER if i > 1 else WRAP
    for i, fill in zip(range(3, 7), (PASS_FILL, PART_FILL, FAIL_FILL, SKIP_FILL)):
        for rr in range(first, r + 1): ws.cell(rr, i).fill = fill

    def block(title, lines):
        nonlocal r
        r += 2
        ws.cell(r, 1, title).font = Font(size=13, bold=True, color='3B1F2B')
        for ln in lines:
            r += 1
            c = ws.cell(r, 1, ln); c.alignment = WRAP
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
            ws.row_dimensions[r].height = 15 * max(1, math.ceil(len(ln) / 105)) + 3

    block('Before you start', [
        '1. Stop the server, deploy the v0.136 build, start the server.',
        '2. Run  .beelz admin reseed replace CONFIRM  once (it keeps a backup). Without it, the preset values are '
        'NOT copied onto abilities your server already has.',
        '3. Check the data is live:  .beelz admin ability AB_Lucie_LiquidFire_AbilityGroup  shows cooldown=10 and freelymove=0.5.',
        '4. Turn on VerboseLogging so the [Beelz CD] / [Beelz TUNE] log lines show what fired.',
    ])
    block('How to test a row', [
        '• "How to test" gives the exact commands. <you> is your character name; <slot> is the spell slot to put it on. '
        'The unit GUID in the give command is only a source label.',
        '• Compare what happens with "Expected result", write what you saw in "Actual result", then pick Worked? '
        '(Pass / Partial / Fail / Can\'t test). The Progress table above updates itself.',
        '• Priority "High" = the hand-picked checklist rows plus every ability whose cooldown Beelzebub now starts '
        'itself (new, untested in game). Filter the Priority column to do those first.',
        '• If a preset seems to cause a problem:  .beelz admin ability <prefab> defaults  removes that one ability\'s preset.',
    ])
    block('What the modified settings mean', [
        '• freelymove = seconds into the cast after which you can move (fixes "locked in place"). If it doesn\'t unroot, '
        'look for freelymove=SKIP(shared) in the log — the edit was refused because parts are shared with another ability.',
        '• cooldown = seconds before you can recast (player only; the boss keeps its own). cooldownscale = multiplier on the native cooldown.',
        '• damagescale = damage multiplier (×0.5 weaker, ×1.5 stronger). leapheight = caps how high a leap throws you.',
        '• summoncap / summontimeout = max summons alive / seconds before summons despawn. range = max aim distance. '
        'forcetimeout = forces a stuck-able effect to end.',
        '• "Native" = the ability\'s own values from the game data, for comparison.',
    ])
    block('Sheets', [
        f'• Presets ({len(presets)}) — enabled abilities that ship with a tester-requested fix. The main test.',
        f'• Hard-blocked ({len(hard)}) — crash / character-breaking abilities, locked in code. Confirm they can\'t be used.',
        f'• Soft-disabled ({len(soft)}) — off by default (non-functional or not worth using). Spot-check a few, and use '
        '"Your call" to mark any you think should ship enabled.',
    ])
    ws.sheet_view.showGridLines = False


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--force', action='store_true', help='overwrite an existing workbook (loses filled-in results)')
    ap.add_argument('--out', default=OUT)
    a = ap.parse_args()
    if os.path.exists(a.out) and not a.force:
        print(f'{a.out} exists (it may hold your results) — pass --force or --out <new path>.'); return
    presets, hard, soft = build()
    wb = openpyxl.Workbook()
    overview(wb, presets, hard, soft)
    sheet(wb, 'Presets', [
        ('#', 5, lambda r: presets.index(r) + 1, 'center'),
        ('Priority', 9, lambda r: r['prio'], 'prio'),
        ('Change type', 14, lambda r: r['types'], 'center'),
        ('Ability', 26, lambda r: f"{r['disp']}\n{r['name']}", 'key'),
        ('Source unit', 18, lambda r: r['unit'], 'text'),
        ('What was modified (setting = value)', 22, lambda r: r['knobs'], 'mono'),
        ('Native (game) values', 14, lambda r: r['native'], 'muted'),
        ('Tester feedback', 46, lambda r: r['tester'], 'text'),
        ('Expected result', 46, lambda r: r['exp'], 'text'),
        ('How to test', 34, lambda r: r['cmd'], 'mono'),
        ('Actual result (your notes)', 40, lambda r: None, 'input'),
        ('Worked?', 12, lambda r: None, 'result'),
    ], presets, '9A1B3A', 'Presets')
    sheet(wb, 'Hard-blocked', [
        ('#', 5, lambda r: hard.index(r) + 1, 'center'),
        ('Ability', 30, lambda r: f"{r['disp']}\n{r['name']}", 'key'),
        ('Source unit', 20, lambda r: r['unit'], 'text'),
        ('Why it is blocked (tester feedback)', 55, lambda r: r['why'], 'text'),
        ('Expected result', 40, lambda r: r['exp'], 'text'),
        ('How to check', 34, lambda r: r['cmd'], 'mono'),
        ('Actual result (your notes)', 36, lambda r: None, 'input'),
        ('Worked?', 12, lambda r: None, 'result'),
    ], hard, '6E7781', 'HardBlocked')
    sheet(wb, 'Soft-disabled', [
        ('#', 5, lambda r: soft.index(r) + 1, 'center'),
        ('Ability', 30, lambda r: f"{r['disp']}\n{r['name']}", 'key'),
        ('Source unit', 20, lambda r: r['unit'], 'text'),
        ('Why it is off (tester feedback)', 55, lambda r: r['tester'] or r['note'], 'text'),
        ('Preset ready if re-enabled', 20, lambda r: r['ready'], 'mono'),
        ('Expected result', 36, lambda r: r['exp'], 'text'),
        ('Actual result (your notes)', 32, lambda r: None, 'input'),
        ('Worked?', 12, lambda r: None, 'result'),
        ('Your call', 13, lambda r: None, 'choice'),
    ], soft, 'B08800', 'SoftDisabled', extra_dv=['Keep off', 'Re-enable', 'Unsure'])
    wb.save(a.out)
    print(f'WROTE {os.path.normpath(a.out)} — presets {len(presets)}, hard-blocked {len(hard)}, soft-disabled {len(soft)}; '
          f"high priority {sum(r['prio'] == 'High' for r in presets)}")


if __name__ == '__main__':
    main()
